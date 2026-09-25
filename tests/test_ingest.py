import pytest

from ingest import chunk_text, ingest_file, extract, IngestLimitExceeded, CHUNK_WORDS, CHUNK_OVERLAP


def test_chunk_text_basic_sizes():
    text = " ".join(f"word{i}" for i in range(500))
    chunks = chunk_text(text)
    assert len(chunks[0].split()) == CHUNK_WORDS
    # sliding window overlap: word at index (size-overlap) of chunk0
    # should be the first word of chunk1
    assert chunks[0].split()[CHUNK_WORDS - CHUNK_OVERLAP] == chunks[1].split()[0]


def test_chunk_text_empty():
    assert chunk_text("") == []
    assert chunk_text("   ") == []


def test_chunk_text_shorter_than_window():
    text = "just a few words here"
    chunks = chunk_text(text)
    assert chunks == [text]


def test_ingest_file_txt(tmp_path):
    p = tmp_path / "note.txt"
    p.write_text("hello world " * 300, encoding="utf-8")
    chunks = ingest_file(str(p))
    assert len(chunks) > 1
    assert all(c["location"] is None for c in chunks)
    assert all(c["text"] for c in chunks)


def test_ingest_file_enforces_chunk_cap(tmp_path, monkeypatch):
    import ingest
    monkeypatch.setattr(ingest, "MAX_CHUNKS_PER_FILE", 3)
    p = tmp_path / "huge.txt"
    p.write_text(" ".join(f"w{i}" for i in range(2000)), encoding="utf-8")
    with pytest.raises(IngestLimitExceeded):
        ingest.ingest_file(str(p))


def test_unsupported_extension(tmp_path):
    p = tmp_path / "file.exe"
    p.write_bytes(b"MZ\x90\x00")
    with pytest.raises(ValueError):
        extract(str(p))


def test_docx_extraction_and_chunk_location(tmp_path):
    pytest.importorskip("docx")
    from docx import Document
    doc = Document()
    doc.add_paragraph("Apples and oranges. " * 20)
    path = tmp_path / "doc.docx"
    doc.save(str(path))

    chunks = ingest_file(str(path))
    assert chunks
    # docx has no natural page/slide concept -> location is None
    assert all(c["location"] is None for c in chunks)


def test_xlsx_extraction_tags_sheet_name(tmp_path):
    pytest.importorskip("openpyxl")
    import openpyxl
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Data"
    ws1.append(["a", "b"])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["hello"])
    path = tmp_path / "book.xlsx"
    wb.save(str(path))

    chunks = ingest_file(str(path))
    locations = {c["location"] for c in chunks}
    assert "sheet 'Data'" in locations
    assert "sheet 'Notes'" in locations


def test_pdf_extraction_tags_page_number(tmp_path):
    pytest.importorskip("reportlab")
    from reportlab.pdfgen import canvas
    path = tmp_path / "doc.pdf"
    c = canvas.Canvas(str(path))
    c.drawString(100, 750, "First page content about foxes.")
    c.showPage()
    c.drawString(100, 750, "Second page content about hounds.")
    c.showPage()
    c.save()

    chunks = ingest_file(str(path))
    locations = [c["location"] for c in chunks]
    assert "p. 1" in locations
    assert "p. 2" in locations


def test_pdf_ocr_only_hits_sparse_pages(monkeypatch):
    """A document with one good text page and one sparse/undecoded page
    should only send the sparse page through OCR, not the whole doc."""
    import ingest

    def fake_text_layer(path):
        return ["Plenty of real text here, well over twenty characters.", ""]

    calls = {}

    def fake_ocr_pages(path, page_numbers):
        calls["page_numbers"] = page_numbers
        return {i: "OCR recovered text" for i in page_numbers}

    monkeypatch.setattr(ingest, "_pdf_text_layer", fake_text_layer)
    monkeypatch.setattr(ingest, "_pdf_ocr_pages", fake_ocr_pages)

    result = ingest._extract_pdf("fake.pdf")

    assert calls["page_numbers"] == [1]   # only the sparse (0-indexed) page
    assert result[0] == ("p. 1", "Plenty of real text here, well over twenty characters.")
    assert result[1] == ("p. 2", "OCR recovered text")
